#!/usr/bin/env python3
"""
Пакетный подбор кодов NTIN по списку названий товаров (Excel или CSV).

  python match_names.py товары.xlsx --out результат.xlsx            # по словам названия
  python match_names.py товары.xlsx --out результат.xlsx --ai       # ИИ сначала разбирает названия (точнее, копейки)
  python match_names.py товары.xlsx --limit 20                      # пробный запуск на первых 20 строках

Для каждой строки: если указан штрихкод, ищем карточку по нему; иначе (или если по штрихкоду пусто) ищем по названию
во всём каталоге. Результат: файл Excel, где у каждого товара есть статус, предлагаемый код, название карточки в каталоге,
пометка «из нашей базы», процент совпадения слов и другие варианты. Ничего не регистрируется и не отправляется в каталог.

Таблица может быть с любыми заголовками; программа ищет колонки «Название», «Бренд», «Штрихкод», «Характеристики».
Найденные ответы каталога запоминаются в match_cache.json: при повторном запуске уже найденное не запрашивается заново.
"""
import argparse
import json
import os
import re
import sys
import time

import pandas as pd

import ai_add
import check_ntin as ck

NAME_HDR = ("название", "наименование", "товар", "name", "product")
BRAND_HDR = ("бренд", "марка", "brand")
GTIN_HDR = ("штрихкод", "штрих-код", "штрих код", "gtin", "ean", "barcode")
NOTE_HDR = ("характеристик", "описание", "размер", "цвет", "примечан", "комментар", "notes")

ST_BARCODE = "Найден по штрихкоду"
ST_NAME = "Найден по названию"
ST_CHECK = "Похожие, проверить"
ST_NONE = "Не найден"
ST_DOWN = "Каталог не ответил"
ACTIONS = {
    ST_BARCODE: "Готово: вносите код",
    ST_NAME: "Почти точно: сверьте объём, размер и цвет",
    ST_CHECK: "Проверить: сверьте название, объём, цвет",
    ST_NONE: "Нужна новая карточка: фото боту «➕ Добавить товар»",
    ST_DOWN: "Запустите ещё раз позже",
}
COLORS = {ST_BARCODE: "C6EFCE", ST_NAME: "C6EFCE", ST_CHECK: "FFEB9C", ST_NONE: "F4B6B6", ST_DOWN: "D9D9D9"}
COLUMNS = ["Статус", "Что делать", "Код NTIN", "Название в каталоге", "Штрихкод в каталоге", "Из нашей базы",
           "Совпадение слов, %", "Другие варианты"]


def gtin_valid(code):
    """Контрольная цифра GTIN-8/12/13/14."""
    if not (code.isascii() and code.isdigit()) or len(code) not in (8, 12, 13, 14):
        return False
    d = [int(x) for x in code]
    return (10 - (sum(d[-2::-2]) * 3 + sum(d[-3::-2])) % 10) % 10 == d[-1]


def pick_column(columns, headers):
    for c in columns:
        low = str(c).lower().strip()
        if any(h in low for h in headers):
            return c
    return None


def read_table(path):
    if path.lower().endswith((".xlsx", ".xlsm", ".xls")):
        df = pd.read_excel(path, dtype=str)
    else:
        df = pd.read_csv(path, dtype=str, sep=None, engine="python", encoding="utf-8-sig")
    df = df.fillna("")
    cols = list(df.columns)
    name = pick_column(cols, NAME_HDR) or cols[0]
    spec = {"name": name,
            "brand": pick_column([c for c in cols if c != name], BRAND_HDR),
            "gtin": pick_column([c for c in cols if c != name], GTIN_HDR),
            "notes": pick_column([c for c in cols if c != name], NOTE_HDR)}
    df = df[~df[name].str.strip().str.lower().str.startswith("пример")]          # строка-пример из шаблона
    df = df[df[name].str.strip() != ""].reset_index(drop=True)
    return df, spec


class Searcher:
    """Поиск по названию с запоминанием ответов и паузой между настоящими запросами."""

    def __init__(self, cache_path=None, pause=0.6, sleep=None, search=None):
        self.path, self.pause, self.sleep = cache_path, pause, sleep or time.sleep
        self.search_fn = search or (lambda q: ck.search_by_name(q, limit=20))
        self.cache, self.new, self.failed = {}, 0, False
        if cache_path and os.path.exists(cache_path):
            try:
                with open(cache_path, encoding="utf-8") as f:
                    self.cache = json.load(f)
            except ValueError:
                self.cache = {}

    def __call__(self, q):
        key = " ".join(ai_add.norm_words(q))
        if key in self.cache:
            return self.cache[key]
        if self.new:
            self.sleep(self.pause)
        cards = self.search_fn(q)
        if cards is None:
            self.failed = True
            return None
        slim = [{"ntin_code": c.get("ntin_code"), "name_ru": c.get("name_ru"), "gtin": c.get("gtin"),
                 "ntin_isdeactivated": bool(c.get("ntin_isdeactivated"))} for c in cards]
        self.cache[key] = slim
        self.new += 1
        if self.new % 25 == 0:
            self.save()
        return slim

    def save(self):
        if self.path:
            with open(self.path, "w", encoding="utf-8") as f:
                json.dump(self.cache, f, ensure_ascii=False)


def heuristic_product(name, brand, notes):
    queries = [name] + ([f"{brand} {name}"] if brand and brand.lower() not in name.lower() else [])
    return ai_add.normalize_product({"kind": name, "brand": brand or None, "features": notes or None, "search_queries": queries})


def normalize_batch(env, rows):
    """ИИ разбирает до ~15 названий за один запрос. rows: список (название, бренд, примечание). Возвращает (товары, расход)."""
    lines = "\n".join(f"{i}. {n}" + (f" | бренд: {b}" if b else "") + (f" | {x}" if x else "") for i, (n, b, x) in enumerate(rows, 1))
    ask = (f"Разбери названия товаров из прайс-листа магазина.\n{lines}\n"
           "Верни только JSON вида {\"items\": [{\"n\": номер, \"kind\": вид товара по-русски 1-3 слова, \"brand\": бренд или null, "
           "\"model_or_article\": модель/артикул/оттенок или null, \"pack\": объём, вес или количество или null, "
           "\"search_queries\": [2-3 коротких запроса по 2-5 слов для поиска такой карточки в каталоге]}]}. "
           "Ничего не выдумывай, чего нет в названии, ставь null. Порядок и число элементов как во входе.")
    text, usage = ai_add.call_ai(env, "Ты помощник магазина в Казахстане. Отвечай только JSON.", [{"type": "text", "text": ask}], max_tokens=4000)
    items = ai_add.parse_json(text).get("items") or []
    by_n = {int(it.get("n")): it for it in items if str(it.get("n", "")).isdigit()}
    out = []
    for i, (n, b, x) in enumerate(rows, 1):
        it = by_n.get(i)
        if it:
            it.setdefault("features", x or None)
            if b and not it.get("brand"):
                it["brand"] = b
            p = ai_add.normalize_product(it)
            p["kind"] = p["kind"] or n
            if not p["search_queries"]:
                p["search_queries"] = [n]
        else:
            p = heuristic_product(n, b, x)
        out.append(p)
    return out, usage


def match_row(product, gtin, search, fetch, base, min_score):
    """Результат для одной строки: словарь с колонками отчёта."""
    res = {"Статус": ST_NONE, "Код NTIN": "", "Название в каталоге": "", "Штрихкод в каталоге": "", "Из нашей базы": "",
           "Совпадение слов, %": "", "Другие варианты": ""}
    search.failed = False
    cands = []
    g = re.sub(r"\D", "", gtin or "")
    if g and gtin_valid(g) and not ck.is_internal(g):
        items = fetch([g])
        if items is None:
            search.failed = True
        for it in items or []:
            n = re.sub(r"\D", "", str(it.get("ntin_code") or ""))
            if n and not it.get("ntin_isdeactivated") and all(n != c["ntin"] for c in cands):
                cands.append({"ntin": n, "name": it.get("name_ru") or "", "gtin": g, "score": 1.0})
        if cands:
            res["Статус"] = ST_BARCODE
            res["Совпадение слов, %"] = 100 if len(cands) == 1 else 90
    if not cands:
        cands = ai_add.find_candidates(product, search, min_score=min_score, limit=6, pause=0)
        if cands:
            top = cands[0]["score"]
            res["Статус"] = ST_NAME if top >= 0.85 else ST_CHECK
            res["Совпадение слов, %"] = round(top * 100)
    if cands:
        for c in cands:
            c["ours"] = c["ntin"] in base
        cands.sort(key=lambda c: (-round(c["score"], 1), not c["ours"], len(c["name"])))      # при равенстве совпадения наши карточки выше
        top = cands[0]
        res.update({"Код NTIN": top["ntin"], "Название в каталоге": top["name"], "Штрихкод в каталоге": top["gtin"],
                    "Из нашей базы": "да" if top["ours"] else "нет",
                    "Другие варианты": "; ".join(f"{c['ntin']} {c['name'][:70]}" for c in cands[1:4])})
    elif search.failed:
        res["Статус"] = ST_DOWN
    res["Что делать"] = ACTIONS[res["Статус"]]
    return res


def write_report(df, path):
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    df.to_excel(path, index=False, sheet_name="Результат")
    wb = load_workbook(path)
    ws = wb["Результат"]
    st_col = list(df.columns).index("Статус") + 1
    for row in ws.iter_rows():
        for cell in row:
            cell.font = Font(name="Arial", size=10, bold=(cell.row == 1))
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for row in range(2, ws.max_row + 1):
        color = COLORS.get(ws.cell(row, st_col).value)
        if color:
            for col in range(st_col, st_col + 2):
                ws.cell(row, col).fill = PatternFill("solid", fgColor=color)
    for i, col in enumerate(df.columns, 1):
        longest = max([len(str(col))] + [len(str(v)) for v in df[col].head(200)])
        ws.column_dimensions[ws.cell(1, i).column_letter].width = min(max(12, longest + 2), 60)
    ws.freeze_panes = "A2"
    wb.save(path)


def run(path, out, env=None, use_ai=False, limit=None, base_path=ck.DEFAULT_BASE, cache_path="match_cache.json",
        search=None, fetch=None, sleep=None):
    df, cols = read_table(path)
    if limit:
        df = df.head(limit)
    base_df = ck.load_base(base_path)
    base = {re.sub(r"\D", "", str(n)) for n in base_df["ntin"]} if len(base_df) else set()
    searcher = Searcher(cache_path, sleep=sleep, search=search)
    fetch = fetch or ck.fetch
    rows = [(r[cols["name"]].strip(), (r[cols["brand"]].strip() if cols["brand"] else ""),
             (r[cols["notes"]].strip() if cols["notes"] else "")) for _, r in df.iterrows()]
    gtins = [(r[cols["gtin"]] if cols["gtin"] else "") for _, r in df.iterrows()]
    products, cost = [], 0.0
    if use_ai:
        for i in range(0, len(rows), 15):
            chunk = rows[i:i + 15]
            try:
                part, usage = normalize_batch(env or {}, chunk)
                cost += usage["cost"]
            except ai_add.AiError as e:
                print(f"ИИ не разобрал строки {i + 1}-{i + len(chunk)} ({e}); для них ищу по словам названия.", file=sys.stderr)
                part = [heuristic_product(*r) for r in chunk]
            products += part
    else:
        products = [heuristic_product(*r) for r in rows]
    min_score = 0.6 if use_ai else 0.5
    results = []
    for i, (product, g) in enumerate(zip(products, gtins), 1):
        results.append(match_row(product, g, searcher, fetch, base, min_score))
        if i % 10 == 0 or i == len(products):
            print(f"  обработано {i} из {len(products)}", flush=True)
    searcher.save()
    report = pd.concat([df.reset_index(drop=True), pd.DataFrame(results)[COLUMNS]], axis=1)
    write_report(report, out)
    counts = report["Статус"].value_counts()
    print(f"\nГотово: {out}")
    for st in (ST_BARCODE, ST_NAME, ST_CHECK, ST_NONE, ST_DOWN):
        if counts.get(st):
            print(f"  {st}: {counts[st]}")
    if use_ai:
        print(f"Расход ИИ примерно ${cost:.4f}")
    if counts.get(ST_DOWN):
        print("Часть товаров не проверена (каталог не ответил): запустите команду ещё раз, найденное запомнено.")
    return report


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("file", help="Excel или CSV со списком товаров")
    ap.add_argument("--out", default="подобранные_коды.xlsx")
    ap.add_argument("--ai", action="store_true", help="ИИ разбирает названия перед поиском (нужен ANTHROPIC_API_KEY в .env)")
    ap.add_argument("--limit", type=int, help="обработать только первые N строк (проба)")
    ap.add_argument("--base", default=ck.DEFAULT_BASE)
    a = ap.parse_args(argv)
    env = None
    if a.ai:
        from nkt_client import _read_dotenv
        env = _read_dotenv(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".env"))
        if not env.get("ANTHROPIC_API_KEY"):
            sys.exit("В .env нет ANTHROPIC_API_KEY: запустите без --ai или добавьте ключ.")
    run(a.file, a.out, env=env, use_ai=a.ai, limit=a.limit, base_path=a.base)


if __name__ == "__main__":
    main()
